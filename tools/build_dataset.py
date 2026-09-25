#!/usr/bin/env python3
"""Build the versioned reference dataset from pinned public inputs."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from pathlib import Path
from urllib.request import urlopen

from bs4 import BeautifulSoup
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
PUBLIC_INPUT = ROOT / "inputs" / "criteria-public.json"
SNAPSHOT_DATE = "2026-09-25"
SOURCE_FILES = {
    "wcag22.json": (
        "https://www.w3.org/WAI/WCAG22/wcag.json",
        "3a034865a879a7d874b60ff2aec7f430cf37bc1597a1e00985680fd80ef5cf75",
    ),
    "wcag21.json": (
        "https://www.w3.org/WAI/WCAG21/wcag.json",
        "fc9f2e8a94421fc305dced3ec4020fde088572493d44da70789ceb6f775bb4d7",
    ),
    "wcag22-en.html": (
        "https://www.w3.org/TR/2024/REC-WCAG22-20241212/",
        "6e3c5fe397257cae509a2fb4752b73062cf8cbeb92c2cec618989b17e4cf7057",
    ),
    "wcag22-ja.html": (
        "https://waic.jp/translations/WCAG22/",
        "a1d56070bfd9b98f5183b7db8dd08082e274305b52572386bb42bc70cba3466b",
    ),
}
WORKBOOK_HASH = "dde704f349397cb224a33d5bf703a73d6318d2ce209d828edc7a643671c20d4b"
CONFORMANCE_IDS = {
    "5": "conformance",
    "5.1": "interpreting-normative-requirements",
    "5.2": "conformance-reqs",
    "5.2.1": "cc1",
    "5.2.2": "cc2",
    "5.2.3": "cc3",
    "5.2.4": "cc4",
    "5.2.5": "cc5",
    "5.3": "conformance-claims",
    "5.3.1": "conformance-required",
    "5.3.2": "conformance-optional",
    "5.4": "conformance-partial",
    "5.5": "conformance-partial-lang",
}
WORKBOOK_COLUMNS = [
    ("english_workbook", "en", "wcag:2.2", "workbook_label"),
    ("jis_2016", "ja", "jis-x-8341-3:2016", "workbook_label"),
    ("wcag20_initial", "ja", "wcag:2.0", "historical_translation"),
    ("wcag20_revised", "ja", "wcag:2.0", "historical_translation"),
    ("wcag21_translation", "ja", "wcag:2.1", "historical_translation"),
]


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def fetch_sources(directory: Path) -> None:
    directory.mkdir(parents=True, exist_ok=True)
    for filename, (url, expected_hash) in SOURCE_FILES.items():
        target = directory / filename
        if target.exists():
            continue
        with urlopen(url, timeout=40) as response:
            data = response.read(2_000_001)
        if len(data) > 2_000_000 or digest(data) != expected_hash:
            raise ValueError(f"Source changed or exceeded limit: {url}")
        target.write_bytes(data)


def source_bytes(directory: Path, filename: str) -> bytes:
    data = (directory / filename).read_bytes()
    expected_hash = SOURCE_FILES[filename][1]
    if digest(data) != expected_hash:
        raise ValueError(f"Checksum mismatch: {filename}; update sources deliberately")
    return data


def normalized_id(value: object) -> str:
    if isinstance(value, (int, float)):
        value = str(value)
        if value.endswith(".0"):
            value = value[:-2]
    if not isinstance(value, str) or not re.fullmatch(r"[1-5](?:\.\d+){0,2}", value):
        raise ValueError(f"Unexpected ID: {value!r}")
    return value


def key(edition: str, number: str) -> str:
    return f"{edition}:{number}"


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_catalog(path: Path, rows: list[dict], items: list[dict], names: list[dict]) -> None:
    by_key = {item["key"]: item for item in items}
    name_index = {(name["item_key"], name.get("workbook_column"), name["status"]): name["value"]
                  for name in names}
    fields = ["number", "kind", "wcag22_status", "level", "wcag22_en", "wcag22_ja",
              "jis2016_ja", "wcag20_initial_ja", "wcag20_revised_ja", "wcag21_ja",
              "wcag22_url", "waic_url", "workbook_row"]
    with path.open("w", encoding="utf-8", newline="") as output:
        writer = csv.DictWriter(output, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        for row in rows:
            number = row["number"]
            item_key = key("wcag:2.2", number)
            item = by_key[item_key]
            official_en = name_index[(item_key, None, "official")]
            official_ja = name_index[(item_key, None, "reference_translation")]
            writer.writerow({
                "number": number, "kind": item["kind"], "wcag22_status": item["status"],
                "level": item["level"] or "", "wcag22_en": official_en,
                "wcag22_ja": official_ja,
                "jis2016_ja": row["labels"][1] if row["labels"][1] not in (None, "-") else "",
                "wcag20_initial_ja": row["labels"][2] if row["labels"][2] not in (None, "-") else "",
                "wcag20_revised_ja": row["labels"][3] if row["labels"][3] not in (None, "-") else "",
                "wcag21_ja": row["labels"][4] if row["labels"][4] not in (None, "-") else "",
                "wcag22_url": f"{SOURCE_FILES['wcag22-en.html'][0]}#{item['anchor']}",
                "waic_url": f"{SOURCE_FILES['wcag22-ja.html'][0]}#{item['anchor']}",
                "workbook_row": row["row"],
            })


def section_name(section) -> str:
    heading = section.find(["h2", "h3", "h4"], recursive=True)
    if heading is None:
        raise ValueError(f"Missing heading: {section.get('id')}")
    parts = [str(x) for x in heading.contents if getattr(x, "name", None) != "bdi"]
    return BeautifulSoup("".join(parts), "lxml").get_text(" ", strip=True)


def section_body(section) -> str:
    excluded = {"header-wrapper", "doclinks", "conformance-level"}
    children = []
    for child in section.find_all(recursive=False):
        if child.name == "section" or excluded.intersection(child.get("class", [])):
            continue
        children.append(str(child))
    return "\n".join(children)


def w3c_tree(data: dict) -> dict[str, tuple[str, str | None, dict]]:
    tree = {}
    for principle in data["principles"]:
        pnum = principle["num"]
        tree[pnum] = ("principle", None, principle)
        for guideline in principle["guidelines"]:
            gnum = guideline["num"]
            tree[gnum] = ("guideline", pnum, guideline)
            for criterion in guideline["successcriteria"]:
                tree[criterion["num"]] = ("success_criterion", gnum, criterion)
    return tree


def workbook_rows(path: Path) -> list[dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook.active
    rows = []
    for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
        if number == 1 or values[0] is None:
            continue
        rows.append({"row": number, "number": normalized_id(values[0]),
                     "raw_number": values[0], "labels": list(values[1:6])})
    if len(rows) != 117 or len({r["number"] for r in rows}) != 117:
        raise ValueError("Workbook must contain 117 unique reference rows")
    return rows


def export_public_input(workbook_path: Path, target: Path) -> None:
    if digest(workbook_path.read_bytes()) != WORKBOOK_HASH:
        raise ValueError("criteria.xlsx differs from the reviewed copy")
    rows = workbook_rows(workbook_path)
    write_json(target, {"source_workbook_sha256": WORKBOOK_HASH,
                        "columns": [column[0] for column in WORKBOOK_COLUMNS],
                        "rows": rows})


def public_rows(path: Path) -> list[dict]:
    seed = json.loads(path.read_text(encoding="utf-8"))
    if set(seed) != {"source_workbook_sha256", "columns", "rows"}:
        raise ValueError("Public input contains unexpected fields")
    if seed["source_workbook_sha256"] != WORKBOOK_HASH:
        raise ValueError("Public input source hash differs")
    if seed["columns"] != [column[0] for column in WORKBOOK_COLUMNS]:
        raise ValueError("Public input columns differ")
    rows = seed["rows"]
    if len(rows) != 117 or len({row["number"] for row in rows}) != 117:
        raise ValueError("Public input must contain 117 unique rows")
    if any(set(row) != {"row", "number", "raw_number", "labels"} or
           len(row["labels"]) != len(WORKBOOK_COLUMNS) or
           normalized_id(row["raw_number"]) != row["number"] for row in rows):
        raise ValueError("Public input row structure differs")
    return rows


def build(source_dir: Path, input_path: Path = PUBLIC_INPUT,
          output_dir: Path | None = None) -> None:
    rows = public_rows(input_path)
    wcag22 = json.loads(source_bytes(source_dir, "wcag22.json"))
    wcag21 = json.loads(source_bytes(source_dir, "wcag21.json"))
    en = BeautifulSoup(source_bytes(source_dir, "wcag22-en.html"), "lxml")
    ja = BeautifulSoup(source_bytes(source_dir, "wcag22-ja.html"), "lxml")
    tree22, tree21 = w3c_tree(wcag22), w3c_tree(wcag21)
    if set(tree22) != {r["number"] for r in rows if not r["number"].startswith("5")}:
        raise ValueError("Workbook and official WCAG 2.2 structure differ")

    sources = [
        {"id": "w3c-wcag22-json", "title": "WCAG 2.2 JSON", "publisher": "W3C",
         "url": SOURCE_FILES["wcag22.json"][0], "sha256": SOURCE_FILES["wcag22.json"][1],
         "retrieved": "2026-09-25", "last_modified": "2026-09-03T15:17:58Z",
         "license_url": "https://github.com/w3c/wcag/tree/main/11ty/json#permission-to-use-with-attribution"},
        {"id": "w3c-wcag21-json", "title": "WCAG 2.1 JSON", "publisher": "W3C",
         "url": SOURCE_FILES["wcag21.json"][0], "sha256": SOURCE_FILES["wcag21.json"][1],
         "retrieved": "2026-09-25", "last_modified": "2026-09-03T15:25:16Z",
         "license_url": "https://github.com/w3c/wcag/tree/main/11ty/json#permission-to-use-with-attribution"},
        {"id": "w3c-wcag22-rec", "title": "WCAG 2.2 Recommendation, 12 December 2024",
         "publisher": "W3C", "url": SOURCE_FILES["wcag22-en.html"][0],
         "sha256": SOURCE_FILES["wcag22-en.html"][1], "retrieved": "2026-09-25",
         "license_url": "https://www.w3.org/copyright/document-license/"},
        {"id": "waic-wcag22-ja", "title": "WCAG 2.2 日本語訳", "publisher": "WAIC翻訳ワーキンググループ",
         "url": SOURCE_FILES["wcag22-ja.html"][0],
         "sha256": SOURCE_FILES["wcag22-ja.html"][1], "retrieved": "2026-09-25",
         "last_modified": "2026-06-08T00:47:03Z",
         "license_url": "https://waic.jp/license-for-translated-documents/",
         "status": "参考訳。正式版はW3Cの英語版"},
        {"id": "criteria-workbook", "title": "達成基準名称集（手元の入力資料）",
         "publisher": "未確認", "file_name": "criteria.xlsx", "sha256": WORKBOOK_HASH,
         "retrieved": "2026-09-25", "publication_status": "private_input; provenance_unverified",
         "public_extract_file": "inputs/criteria-public.json",
         "public_extract_sha256": digest(input_path.read_bytes())},
        {"id": "waic-jis2016-guide", "title": "JIS X 8341-3:2016 解説", "publisher": "WAIC",
         "url": "https://waic.jp/docs/jis2016/understanding/",
         "note": "JIS X 8341-3:2016 と WCAG 2.0 の一致規格という関係の根拠"},
    ]

    items, names, texts, relations = [], [], [], []
    for row in rows:
        number = row["number"]
        conformance = number.startswith("5")
        if conformance:
            kind = "conformance_section"
            parent = number.rpartition(".")[0] or None
            anchor = CONFORMANCE_IDS[number]
            versions = ["2.2"]
            level = None
            english_name = section_name(en.find("section", id=anchor))
            english_name = re.sub(r"^5(?:\.\d+)*\.?\s*", "", english_name)
            w3c_source = "w3c-wcag22-rec"
        else:
            kind, parent, node22 = tree22[number]
            anchor = node22["id"]
            versions = node22["versions"]
            level = node22.get("level") or (tree21[number][2].get("level") if number in tree21 else None)
            english_name = node22["handle"]
            w3c_source = "w3c-wcag22-json"
        if number == "4.1.1":
            status = "obsolete_in_2.2"
        else:
            status = "active"
        item22 = key("wcag:2.2", number)
        items.append({"key": item22, "standard": "WCAG", "edition": "2.2", "number": number,
                      "kind": kind, "parent": key("wcag:2.2", parent) if parent else None,
                      "status": status, "level": level if kind == "success_criterion" and status == "active" else None,
                      "anchor": anchor, "source_id": w3c_source,
                      "workbook_row": row["row"], "workbook_number_value": row["raw_number"]})

        if conformance:
            section_en = en.find("section", id=anchor)
            english_body = section_body(section_en)
            english_text_source = "w3c-wcag22-rec"
            english_extract = "section_body_html"
        else:
            english_body = node22["content"]
            english_text_source = "w3c-wcag22-json"
            english_extract = "official_json_content"
        section_ja = ja.find("section", id=anchor)
        if section_ja is None or not english_body:
            raise ValueError(f"Missing text: {number}")
        japanese_name = section_name(section_ja)
        japanese_name = re.sub(r"^5(?:\.\d+)*\.?\s*", "", japanese_name) if conformance else japanese_name
        names.extend([
            {"item_key": item22, "language": "en", "value": english_name,
             "status": "official", "source_id": w3c_source,
             "source_ref": f"{SOURCE_FILES['wcag22-en.html'][0]}#{anchor}"},
            {"item_key": item22, "language": "ja", "value": japanese_name,
             "status": "reference_translation", "source_id": "waic-wcag22-ja",
             "source_ref": f"{SOURCE_FILES['wcag22-ja.html'][0]}#{anchor}"},
        ])
        texts.extend([
            {"item_key": item22, "language": "en", "format": "html_fragment",
             "html": english_body, "extraction": english_extract,
             "source_id": english_text_source,
             "source_ref": f"{SOURCE_FILES['wcag22-en.html'][0]}#{anchor}"},
            {"item_key": item22, "language": "ja", "format": "html_fragment",
             "html": section_body(section_ja), "extraction": "section_body_html",
             "source_id": "waic-wcag22-ja",
             "source_ref": f"{SOURCE_FILES['wcag22-ja.html'][0]}#{anchor}"},
        ])

        # Edition-specific historical items are structural references, not claims
        # that a later revision's wording is identical to an earlier edition.
        if not conformance:
            for edition in ("2.0", "2.1"):
                if edition not in versions:
                    continue
                old_node = tree21[number][2]
                historical_key = key(f"wcag:{edition}", number)
                items.append({"key": historical_key, "standard": "WCAG", "edition": edition,
                              "number": number, "kind": kind,
                              "parent": key(f"wcag:{edition}", parent) if parent else None,
                              "status": "active", "level": old_node.get("level") or None,
                              "anchor": old_node["id"], "source_id": "w3c-wcag21-json",
                              "content_status": "text_recorded" if edition == "2.1" else "not_yet_collected"})
                if edition == "2.1":
                    names.append({"item_key": historical_key, "language": "en",
                                  "value": old_node["handle"], "status": "official",
                                  "source_id": "w3c-wcag21-json",
                                  "source_ref": f"https://www.w3.org/TR/WCAG21/#{old_node['id']}"})
                    texts.append({"item_key": historical_key, "language": "en",
                                  "format": "html_fragment", "html": old_node["content"],
                                  "extraction": "official_json_content", "source_id": "w3c-wcag21-json",
                                  "source_ref": f"https://www.w3.org/TR/WCAG21/#{old_node['id']}"})
            if "2.0" in versions and "2.1" in versions:
                relations.append({"from": key("wcag:2.0", number), "to": key("wcag:2.1", number),
                                  "relation": "same_number_in_later_edition", "source_id": "w3c-wcag21-json"})
            if "2.1" in versions:
                relations.append({"from": key("wcag:2.1", number), "to": item22,
                                  "relation": "removed_in_later_edition" if number == "4.1.1" else "same_number_in_later_edition",
                                  "source_id": "w3c-wcag22-json"})

        for column_index, (column_key, language, scope, name_status) in enumerate(WORKBOOK_COLUMNS):
            value = row["labels"][column_index]
            if value is not None and not isinstance(value, str):
                raise ValueError(f"Unexpected label type at row {row['row']}")
            actual_key = key(scope, number)
            if value not in (None, "-"):
                if actual_key not in {item["key"] for item in items}:
                    if scope == "jis-x-8341-3:2016":
                        items.append({"key": actual_key, "standard": "JIS X 8341-3",
                                      "edition": "2016", "number": number, "kind": kind,
                                      "parent": key(scope, parent) if parent and parent != "5.1" else None,
                                      "status": "in_workbook", "level": tree21[number][2].get("level") if number in tree21 and kind == "success_criterion" else None,
                                      "source_id": "criteria-workbook", "content_status": "not_included"})
                        if number in tree21:
                            relations.append({"from": actual_key, "to": key("wcag:2.0", number),
                                              "relation": "corresponding_number_in_identical_standard",
                                              "source_id": "waic-jis2016-guide"})
                    elif scope == "wcag:2.0" or scope == "wcag:2.1":
                        raise ValueError(f"Historical name without edition item: {actual_key}")
            if actual_key not in {item["key"] for item in items}:
                actual_key = item22
            names.append({"item_key": actual_key, "language": language,
                          "value": value if value not in (None, "-") else None,
                          "status": name_status,
                          "missing_marker": "blank_in_workbook" if value is None else "dash_in_workbook" if value == "-" else None,
                          "workbook_column": column_key,
                          "source_id": "criteria-workbook",
                          "source_cell": f"シート1!{chr(ord('B') + column_index)}{row['row']}"})

    # Check every reference after all items have been collected.
    item_keys = {item["key"] for item in items}
    if len(item_keys) != len(items):
        raise ValueError("Duplicate item keys")
    for item in items:
        if item["parent"] and item["parent"] not in item_keys:
            raise ValueError(f"Missing parent: {item['key']}")
    for relation in relations:
        if relation["from"] not in item_keys or relation["to"] not in item_keys:
            raise ValueError(f"Broken relation: {relation}")
    for record in names + texts:
        if record["item_key"] not in item_keys:
            raise ValueError(f"Broken item reference: {record['item_key']}")

    out = output_dir or ROOT / "data"
    for filename, value in [
        ("sources.json", sources), ("items.json", items), ("names.json", names),
        ("texts.json", texts), ("relations.json", relations),
    ]:
        write_json(out / filename, value)
    write_catalog(out / "catalog.csv", rows, items, names)
    write_json(out / "manifest.json", {
        "dataset_version": "0.2.0", "generated_on": SNAPSHOT_DATE,
        "status": "draft_local; not_reviewed_for_public_release",
        "counts": {"workbook_rows": len(rows), "items": len(items), "names": len(names),
                   "texts": len(texts), "relations": len(relations)},
        "files": ["sources.json", "items.json", "names.json", "texts.json", "relations.json", "catalog.csv"],
    })
    print(f"Built {len(items)} items, {len(names)} names, {len(texts)} texts, {len(relations)} relations")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=ROOT / ".cache")
    parser.add_argument("--input", type=Path, default=PUBLIC_INPUT)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "data")
    parser.add_argument("--export-public-input", action="store_true",
                        help="extract only approved columns from local criteria.xlsx")
    parser.add_argument("--criteria", type=Path, default=ROOT / "criteria.xlsx")
    parser.add_argument("--fetch", action="store_true", help="download pinned official sources into source-dir")
    args = parser.parse_args()
    if args.export_public_input:
        export_public_input(args.criteria, args.input)
        return
    if args.fetch:
        fetch_sources(args.source_dir)
    build(args.source_dir, args.input, args.output_dir)


if __name__ == "__main__":
    main()
